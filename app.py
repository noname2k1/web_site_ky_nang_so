import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

from flask import Flask, jsonify, request, send_from_directory, send_file

from database import (
    delete_result,
    delete_results,
    departments,
    get_result,
    init_db,
    insert_result,
    list_results,
    statistics,
)


BASE_DIR = Path(__file__).resolve().parent


app = Flask(
    __name__,
    static_folder=str(BASE_DIR),
    static_url_path="",
)


# ============================================================
# DATABASE
# ============================================================

try:
    init_db()
    print("DATABASE INITIALIZED SUCCESSFULLY")
except Exception as e:
    print("DATABASE INIT ERROR:", repr(e))


# ============================================================
# HELPERS
# ============================================================


def clean_text(value, max_len=200):
    value = str(value or "").strip()

    value = re.sub(
        r"[\x00-\x1f\x7f]",
        "",
        value,
    )

    return value[:max_len]


def int_value(value, default=0, minimum=0):
    try:
        return max(
            minimum,
            int(value),
        )
    except (TypeError, ValueError):
        return default


def float_value(value, default=0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


# ============================================================
# PAGES
# ============================================================


@app.get("/")
def index():
    return send_from_directory(
        BASE_DIR,
        "index.html",
    )


@app.get("/admin")
def admin():
    return send_from_directory(
        BASE_DIR / "admin",
        "admin.html",
    )


# ============================================================
# SUBMIT RESULT
# ============================================================


@app.post("/api/submit")
def api_submit():

    data = request.get_json(silent=True) or {}

    # --------------------------------------------------------
    # BASIC INFO
    # --------------------------------------------------------

    name = clean_text(
        data.get("name"),
        100,
    )

    department = clean_text(
        data.get("department"),
        100,
    )

    if len(name) < 2:
        return jsonify(
            success=False,
            message="Họ và tên không hợp lệ.",
        ), 400

    if len(department) < 2:
        return jsonify(
            success=False,
            message="Phòng ban không hợp lệ.",
        ), 400

    # --------------------------------------------------------
    # SCORE
    # --------------------------------------------------------

    question_count = int_value(data.get("questionCount"))

    answered_count = int_value(data.get("answeredCount"))

    correct_count = int_value(data.get("correctCount"))

    wrong_count = int_value(data.get("wrongCount"))

    configured_minutes = int_value(data.get("configuredMinutes"))

    remaining_seconds = int_value(data.get("remainingSeconds"))

    percentage = max(
        0,
        min(
            100,
            float_value(data.get("percentage")),
        ),
    )

    auto_submit = bool(data.get("autoSubmit"))

    # --------------------------------------------------------
    # VALIDATE SCORE
    # --------------------------------------------------------

    if question_count < 0:
        question_count = 0

    if correct_count > question_count:
        correct_count = question_count

    if answered_count > question_count:
        answered_count = question_count

    # Sai + bỏ trống
    wrong_count = max(
        0,
        question_count - correct_count,
    )

    # --------------------------------------------------------
    # DETAILS
    # --------------------------------------------------------

    details = data.get(
        "details",
        [],
    )

    if not isinstance(details, list):
        details = []

    safe_details = []

    for item in details:
        if not isinstance(item, dict):
            continue

        safe_details.append(
            {
                "number": int_value(item.get("number")),
                "questionId": int_value(item.get("questionId")),
                "question": clean_text(
                    item.get("question"),
                    2000,
                ),
                "selected": item.get("selected"),
                "selectedLabel": clean_text(
                    item.get("selectedLabel"),
                    20,
                ),
                "correct": item.get("correct"),
                "correctLabel": clean_text(
                    item.get("correctLabel"),
                    20,
                ),
                "correctText": clean_text(
                    item.get("correctText"),
                    2000,
                ),
                "isCorrect": bool(item.get("isCorrect")),
            }
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    now = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    try:
        row_id = insert_result(
            {
                "submitTime": now,
                "name": name,
                "department": department,
                "questionCount": question_count,
                "configuredMinutes": configured_minutes,
                "remainingSeconds": remaining_seconds,
                "submitType": ("Hết giờ" if auto_submit else "Tự nộp"),
                "answeredCount": answered_count,
                "correctCount": correct_count,
                "wrongCount": wrong_count,
                "percentage": percentage,
                "details": safe_details,
            }
        )

        return jsonify(
            success=True,
            id=row_id,
            message="Đã lưu kết quả.",
        )

    except Exception as e:
        import traceback

        print("========== DATABASE ERROR ==========")

        traceback.print_exc()

        print("====================================")

        return jsonify(
            success=False,
            message=(f"{type(e).__name__}: {e}"),
        ), 500


# ============================================================
# RESULTS
# ============================================================


@app.get("/api/results")
def api_results():

    search = clean_text(
        request.args.get("search"),
        100,
    )

    department = clean_text(
        request.args.get("department"),
        100,
    )

    try:
        results = list_results(
            search,
            department,
        )

        return jsonify(
            success=True,
            results=results,
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# RESULT DETAIL
# ============================================================


@app.get("/api/results/<int:result_id>")
def api_result_detail(result_id):

    try:
        result = get_result(result_id)

        if result is None:
            return jsonify(
                success=False,
                message="Không tìm thấy kết quả.",
            ), 404

        return jsonify(
            success=True,
            result=result,
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# DELETE ONE
# ============================================================


@app.delete("/api/results/<int:result_id>")
def api_delete_result(result_id):

    try:
        deleted = delete_result(result_id)

        if not deleted:
            return jsonify(
                success=False,
                message="Không tìm thấy kết quả.",
            ), 404

        return jsonify(
            success=True,
            message="Đã xóa kết quả.",
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# DELETE MANY
# ============================================================


@app.post("/api/results/delete-many")
def api_delete_many():

    data = request.get_json(silent=True) or {}

    ids = data.get(
        "ids",
        [],
    )

    if not isinstance(ids, list):
        return jsonify(
            success=False,
            message="Danh sách ID không hợp lệ.",
        ), 400

    try:
        count = delete_results(ids)

        return jsonify(
            success=True,
            deleted=count,
        )

    except (TypeError, ValueError) as e:
        return jsonify(
            success=False,
            message=str(e),
        ), 400

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# DEPARTMENTS
# ============================================================


@app.get("/api/departments")
def api_departments():

    try:
        return jsonify(
            success=True,
            departments=departments(),
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# STATISTICS
# ============================================================


@app.get("/api/statistics")
def api_statistics():

    try:
        return jsonify(
            success=True,
            statistics=statistics(),
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# EXPORT EXCEL
# ============================================================


@app.get("/api/export")
def api_export():

    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Alignment,
            Font,
        )
        from openpyxl.utils import (
            get_column_letter,
        )

    except ImportError:
        return jsonify(
            success=False,
            message=("Chưa cài openpyxl. Chạy: pip install openpyxl"),
        ), 500

    try:
        rows = list_results(
            clean_text(
                request.args.get("search"),
                100,
            ),
            clean_text(
                request.args.get("department"),
                100,
            ),
        )

        wb = Workbook()

        ws = wb.active

        ws.title = "Kết quả"

        headers = [
            "STT",
            "Thời gian nộp",
            "Họ và tên",
            "Phòng ban",
            "Số câu",
            "Thời gian quy định (phút)",
            "Thời gian còn lại (giây)",
            "Hình thức",
            "Đã trả lời",
            "Đúng",
            "Sai/Bỏ trống",
            "Điểm (%)",
        ]

        ws.append(headers)

        # ----------------------------------------------------
        # HEADER
        # ----------------------------------------------------

        for cell in ws[1]:
            cell.font = Font(bold=True)

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # ----------------------------------------------------
        # DATA
        # ----------------------------------------------------

        for i, row in enumerate(
            rows,
            1,
        ):
            ws.append(
                [
                    i,
                    row["submit_time"],
                    row["name"],
                    row["department"],
                    row["question_count"],
                    (
                        row["configured_minutes"]
                        if row["configured_minutes"]
                        else "Không giới hạn"
                    ),
                    row["remaining_seconds"],
                    row["submit_type"],
                    row["answered_count"],
                    row["correct_count"],
                    row["wrong_count"],
                    row["percentage"],
                ]
            )

        # ----------------------------------------------------
        # WIDTH
        # ----------------------------------------------------

        widths = [
            8,
            22,
            28,
            28,
            10,
            24,
            24,
            16,
            14,
            10,
            16,
            12,
        ]

        for i, width in enumerate(
            widths,
            1,
        ):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions

        # ----------------------------------------------------
        # SAVE TO MEMORY
        # ----------------------------------------------------

        output = BytesIO()

        wb.save(output)

        output.seek(0)

        filename = "ket_qua_thi.xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )

    except Exception as e:
        import traceback

        traceback.print_exc()

        return jsonify(
            success=False,
            message=f"{type(e).__name__}: {e}",
        ), 500


# ============================================================
# LOCAL DEVELOPMENT
# ============================================================

if __name__ == "__main__":
    print("=" * 60)

    print("HE THONG THI TRAC NGHIEM")

    print("Trang thi : http://127.0.0.1:5000")

    print("Quan ly   : http://127.0.0.1:5000/admin")

    print("=" * 60)

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False,
    )
